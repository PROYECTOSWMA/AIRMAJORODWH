"""Tests del parser del reporte de ventas estático (``parse_sales_selva``).

Se cubren las irregularidades reales del libro "1.- REPORTE DE VENTAS": hojas con
cabeceras distintas, columnas sin título (la FECHA de la columna A y columnas
numéricas sueltas), encabezados repetidos (``CREDITOS``), el marcador de "sin dato"
``"-"`` y las filas vacías de separación. Además se corre sobre el archivo real del
repo para fijar el contrato del asset ``static_data_asset``.
"""
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from assets.static_data_asset.parsers.sales_selva import (
    _column_names,
    _clean,
    parse_sales_selva,
)

_REPO_ROOT = Path(__file__).resolve().parents[2]
_VENTAS_PATH = _REPO_ROOT / "static" / "1.- REPORTE DE VENTAS AL 19.092026.xlsx"


def _libro(path: Path) -> None:
    """Libro sintético con las variantes de cabecera que hay que soportar."""
    wb = openpyxl.Workbook()
    selva = wb.active
    selva.title = "SELVA X-26"
    # Columna A sin cabecera (es FECHA), AVION y una columna numérica sin nombre.
    selva.append(["", "NRO MANIF", "AVION", "", "TOTAL"])
    selva.append([datetime(2026, 9, 1), "006-1", "OB-1", 5, 100])
    selva.append(["", "", "", "", ""])  # fila vacía -> se descarta
    selva.append(["-", "-", "-", "-", "-"])  # solo "-" -> se descarta
    selva.append([datetime(2026, 9, 2), "006-2", "OB-2", None, "-"])

    sa = wb.create_sheet("S.A Y-26")
    # CREDITOS repetido: el segundo se sufija `_2`.
    sa.append(["FECHA", "NRO", "CREDITOS", "CONTADO", "CREDITOS"])
    sa.append([datetime(2026, 9, 3), "006-3", "a", "b", "c"])

    wb.save(path)


def _libro_matriculas(path: Path) -> None:
    """Hoja con la columna AVION llena de tecleos, como el libro real."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SELVA F-26"
    ws.append(["FECHA", "NRO", "AVION", "TOTAL"])
    ws.append([datetime(2026, 2, 18), "007-1733", "0B-2180", 1660])  # el cero por la O
    ws.append([datetime(2026, 2, 20), "006-4914", "ob.2209", 4500])  # minúsculas y punto
    ws.append([datetime(2026, 2, 28), "008-2784", "OB2221", 9350])  # sin separador
    ws.append([datetime(2026, 1, 6), "003-4560", "003-4560", 3900])  # nº de documento
    ws.append([datetime(2026, 1, 7), "001-5848", "SAETA", 6745])  # otra aerolínea
    wb.save(path)


def test_parse_sales_selva_normaliza_las_matriculas(tmp_path):
    """La matrícula queda ``'OB-1234'`` (y vacía si la celda trae un nº de documento).

    De esta columna sale el centro de costo en dbt (``extract_cost_center``), así que sin
    normalizar un ``'0B-2180'`` producía el CC basura ``'0B-2'``.
    """
    path = tmp_path / "ventas.xlsx"
    _libro_matriculas(path)

    filas = parse_sales_selva(path)

    assert [f.get("avion") for f in filas] == [
        "OB-2180",
        "OB-2209",
        "OB-2221",
        None,
        "SAETA",
    ]


def test_column_names_resuelve_vacios_y_repetidos():
    assert _column_names(("", "NRO MANIF", "AVION", "", "CREDITOS", "CREDITOS")) == [
        "fecha",
        "nro_manif",
        "avion",
        "col_3",
        "creditos",
        "creditos_2",
    ]


def test_clean_convierte_vacios_y_guion_a_none():
    assert _clean(None) is None
    assert _clean("   ") is None
    assert _clean("-") is None
    assert _clean(" 010-1 ") == "010-1"
    assert _clean(12) == 12


def test_parse_sales_selva_sobre_libro_sintetico(tmp_path):
    path = tmp_path / "ventas.xlsx"
    _libro(path)

    filas = parse_sales_selva(path)

    # 2 filas de la primera hoja + 1 de la segunda (las vacías/'-' se descartan).
    assert len(filas) == 3
    assert all("sheet_name" in f and "row_number" in f for f in filas)

    selva = filas[0]
    assert selva["sheet_name"] == "SELVA X-26"
    assert selva["row_number"] == 2
    assert selva["fecha"] == datetime(2026, 9, 1)
    assert selva["nro_manif"] == "006-1"
    assert selva["avion"] == "OB-1"
    assert selva["col_3"] == 5
    assert selva["total"] == 100

    # Segunda fila de SELVA: las celdas "-"/vacías quedan como None (clave ausente).
    segunda = filas[1]
    assert segunda["row_number"] == 5
    assert segunda["nro_manif"] == "006-2"
    assert "col_3" not in segunda
    assert "total" not in segunda

    # Hoja S.A: encabezado CREDITOS repetido -> creditos / creditos_2.
    sa = filas[2]
    assert sa["sheet_name"] == "S.A Y-26"
    assert sa["fecha"] == datetime(2026, 9, 3)
    assert sa["nro"] == "006-3"
    assert sa["creditos"] == "a"
    assert sa["contado"] == "b"
    assert sa["creditos_2"] == "c"


@pytest.mark.skipif(not _VENTAS_PATH.is_file(), reason="falta el Excel estático")
def test_parse_sales_selva_sobre_archivo_real():
    filas = parse_sales_selva(_VENTAS_PATH)

    # El libro trae ~12 mil filas de datos repartidas en 43 hojas con datos.
    assert len(filas) > 10000
    hojas = {f["sheet_name"] for f in filas}
    assert "SELVA SET-26" in hojas
    assert "S.A SET-26" in hojas
    assert "2026" not in hojas  # hoja vacía, sin filas

    primera = filas[0]
    assert primera["sheet_name"] == "SELVA SET-26"
    assert primera["row_number"] == 2
    assert str(primera["fecha"]).startswith("2026-09-01")
    assert primera["nro_manif"] == "006-5575"
    assert primera["avion"] == "OB-2221"
    assert primera["agente"] == "LINDA"
    assert primera["tipo_de_vuelo"] == "COMERCIAL"
    assert primera["total"] == 2.84

    # Ninguna matrícula se queda con el tecleo del 0 ('0B-2180' -> 'OB-2180') ni con un
    # nº de documento ('003-4560'): de ahí salían los CC basura '0B-2' y '003-'.
    aviones = {f["avion"] for f in filas if f.get("avion")}
    assert not any(str(a).startswith("0") for a in aviones)
    assert "003-4560" not in aviones
    assert "OB-2180" in aviones and "OB-2221" in aviones
