
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from assets.static_data_asset.parsers.ambulance_sales.doc_parser import parse_factura_fields
from assets.utils import normalize_matricula, to_db_field

AMBULANCIA_SHEET_NAME = "VUELOS B200 2026"

# Columnas A..Y (1..25) del enunciado.
AMBULANCIA_MAX_COL = 25

# El CODIGO de la tabla es la FECHA del vuelo en `ddmmaaaa` (8 dígitos): es la clave que
# separa las filas de la tabla del bloque de COSTO por hora que hay DEBAJO (ver docstring).
_CODIGO_TABLA = re.compile(r"^\d{8}$")


def _is_empty(value: Any) -> bool:
    """True si el valor es ``None`` o un texto vacío (celdas vacías del Excel)."""
    return value is None or (isinstance(value, str) and value.strip() == "")


def _texto_codigo(value: Any) -> str:
    """CODIGO de la fila como texto (las celdas numéricas llegan como ``float``)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_ambulancia_ventas(path: Path) -> list[dict[str, Any]]:
    """Lee la hoja "VUELOS B200 2026" y devuelve SOLO las filas de su tabla (A..Y).

    ⚠️ **Debajo de la tabla la hoja trae OTRA tabla** (el COSTO por hora de cada
    aeronave) y sus filas también tienen la columna A con dato (``'HORAS'``, 60, 65, 70,
    75, 80): con el criterio de "columna A no vacía" entraban como si fuesen vuelos —6
    filas basura, con la aeronave y la fecha corridas una columna—. Por eso una fila
    solo entra si su **CODIGO tiene la forma de la tabla**: la fecha del vuelo en
    ``ddmmaaaa`` (8 dígitos). No se corta la lectura por el primer hueco porque
    *dentro* de la tabla los hay de hasta 11 filas (bloques por mes).
    """
    wb = load_workbook(path, data_only=True)
    ws = wb[AMBULANCIA_SHEET_NAME]

    # Columnas A..Y que tienen cabecera (el resto se ignora).
    headers: list[str] = []
    cols: list[int] = []
    for c in range(1, AMBULANCIA_MAX_COL + 1):
        header = ws.cell(row=1, column=c).value
        if _is_empty(header):
            continue
        headers.append(to_db_field(str(header)))
        cols.append(c)

    filas: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        # Solo las filas de la tabla: el bloque de costos de abajo tiene la columna A
        # con 'HORAS' o con números de 2 dígitos (no son fechas ddmmaaaa).
        if not _CODIGO_TABLA.match(_texto_codigo(ws.cell(row=r, column=1).value)):
            continue
        fila: dict[str, Any] = {
            h: ws.cell(row=r, column=c).value for h, c in zip(headers, cols)
        }
        # Matrícula canónica ('OB2209' / '0B2209' -> 'OB-2209'): de ella sale el centro
        # de costo en dbt (`extract_cost_center`).
        if "aeronave" in fila:
            fila["aeronave"] = normalize_matricula(fila["aeronave"])
        # Columna derivada de `factura` (la calcula el parser, no el Excel).
        fila.update(parse_factura_fields(fila.get("factura")))
        filas.append(fila)

    return filas
