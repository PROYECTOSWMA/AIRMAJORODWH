"""Parser del Excel estático "REPORTE DE VENTAS" (hoja por mes y empresa).

El libro ``static/1.- REPORTE DE VENTAS AL 19.092026.xlsx`` trae **una hoja por mes
y empresa** (``SELVA ...`` = Air Majoro Selva, ``S.A ...`` = Air Majoro S.A.), con la
misma idea de reporte de ventas pero con cabeceras que **varían** de una hoja a otra:

* columnas sin título (p. ej. la FECHA de algunas hojas viene en la columna A pero sin
  cabecera; en las hojas ``S.A`` hay una columna numérica sin nombre entre
  ``DEPOSITO`` y ``ESSALUD``/``CREDITOS``);
* encabezados repetidos en la misma hoja (``CREDITOS`` aparece dos veces en las hojas
  ``S.A``);
* sinónimos y erratas (``NRO``/``NRO MANIF``, ``PASAJE``/``PASAJES``,
  ``VUELO``/``TIPO DE VUELO``, ``C0NTADO``/``CONTADO``, ...).

Para no perder información se consolida **todo** con estas reglas:

* la cabecera se normaliza con ``to_db_field``;
* la columna A siempre es ``fecha`` aunque su cabecera venga vacía;
* las columnas sin cabecera se nombran ``col_<índice>`` (índice 0-based de la hoja);
* si un nombre se repite en la misma hoja, el segundo se sufija ``_2`` (y así siguiendo);
* las celdas vacías o con ``"-"`` (marcador de "sin dato" del libro) se guardan como
  ``None``;
* se descartan las filas completamente vacías, incluidas las de separación que solo
  traen ``"-"``;
* la **matrícula** de la aeronave se normaliza a ``'OB-1234'`` (``normalize_matricula``):
  los reportes se teclean a mano (``'0B-2180'``, ``'ob.2209'``, ``'OB2221'``) y de esa
  columna sale el centro de costo en dbt. Si la celda trae un nº de documento en vez
  de la matrícula (``'003-4560'``) se deja vacía, para no inventar un CC.

Cada fila se etiqueta con ``sheet_name`` (nombre de la hoja) y ``row_number`` (fila de
Excel, 1-based; la cabecera es la fila 1) para poder rastrearla. El asset
``static_data_asset`` vuelca el resultado en la tabla ``sales_selva``.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from assets.utils import normalize_matricula, to_db_field

# Valores que el libro usa como "sin dato".
_NULL_MARKERS = frozenset({"", "-"})


def _clean(value: Any) -> Any:
    """Texto con ``strip``; ``""`` y ``"-"`` se convierten a ``None``."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return None if text in _NULL_MARKERS else text
    return value


def _column_names(header: tuple) -> list[str]:
    """Nombre de campo de cada columna de la hoja.

    Resuelve las columnas sin cabecera (``col_<índice>``, salvo la columna A que es
    siempre ``fecha``) y los encabezados repetidos (sufijo ``_2``, ``_3``, ...),
    de modo que cada nombre sea único dentro de la hoja.
    """
    names: list[str] = []
    used: dict[str, int] = {}
    for index, raw in enumerate(header):
        name = to_db_field(str(raw)) if raw is not None and str(raw).strip() else ""
        if not name:
            name = "fecha" if index == 0 else f"col_{index}"
        used[name] = used.get(name, 0) + 1
        names.append(name if used[name] == 1 else f"{name}_{used[name]}")
    return names


def parse_sales_selva(path: Path) -> list[dict[str, Any]]:
    """Lee todas las hojas del libro y devuelve una fila por fila con datos.

    El contrato es el de los demás parsers estáticos: ``parser(path) -> list[dict]``
    con las claves ya normalizadas al modelo ``SalesSelva`` (más ``sheet_name`` y
    ``row_number``). Las filas completamente vacías no generan registro.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    filas: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            names = _column_names(header)
            for row_number, row in enumerate(rows, start=2):
                cells = [_clean(value) for value in row]
                if len(cells) < len(names):
                    cells += [None] * (len(names) - len(cells))
                if not any(value is not None for value in cells):
                    continue
                fila: dict[str, Any] = {
                    name: value
                    for name, value in zip(names, cells)
                    if value is not None
                }
                # La matrícula llega con tecleos ('0B-2180', 'ob.2209', 'OB2221') y de
                # ella sale el centro de costo en dbt: se normaliza a 'OB-1234' (y si la
                # celda trae un nº de documento en vez de la aeronave, se vacía).
                if "avion" in fila:
                    avion = normalize_matricula(fila["avion"])
                    if avion is None:
                        del fila["avion"]
                    else:
                        fila["avion"] = avion
                fila["sheet_name"] = ws.title
                fila["row_number"] = row_number
                filas.append(fila)
    finally:
        wb.close()
    return filas
