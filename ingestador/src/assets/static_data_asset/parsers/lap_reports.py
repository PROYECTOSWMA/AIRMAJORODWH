"""Parser del Excel LAP: hoja "Aeronave".

Toma la primera fila como cabecera, ignora las columnas sin cabecera, descarta las
filas totalmente vacías y devuelve una lista de dicts con las claves normalizadas
(snake_case) para insertar en ``lap_reports``.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from assets.utils import to_db_field

LAP_SHEET_NAME = "Aeronave"


def _is_empty(value: Any) -> bool:
    """True si el valor es ``None`` o un texto vacío (celdas vacías del Excel)."""
    return value is None or (isinstance(value, str) and value.strip() == "")


def parse_lap_aeronave(path: Path) -> list[dict[str, Any]]:
    """Lee la hoja "Aeronave" y devuelve las filas de las columnas con cabecera."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[LAP_SHEET_NAME]
    it = ws.iter_rows(values_only=True)
    header = next(it)

    # Solo las columnas que tienen cabecera no vacía.
    cols = [i for i, h in enumerate(header) if not _is_empty(h)]
    headers = [to_db_field(str(header[i])) for i in cols]

    filas: list[dict[str, Any]] = []
    for row in it:
        if not row or all(_is_empty(v) for v in row):
            continue
        fila: dict[str, Any] = {}
        for i, h in zip(cols, headers):
            v = row[i] if i < len(row) else None
            if _is_empty(v):
                fila[h] = None
            elif isinstance(v, str):
                fila[h] = v.strip()
            else:
                fila[h] = v
        filas.append(fila)

    wb.close()
    return filas
